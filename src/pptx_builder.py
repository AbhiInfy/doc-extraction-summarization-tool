from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .models import DocumentSummary, FeatureSummary, ThemeSummary

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_MID = RGBColor(0x1A, 0x36, 0x5D)
ORACLE_RED = RGBColor(0xC7, 0x46, 0x34)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GOLD = RGBColor(0xB4, 0x53, 0x09)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF6, 0xF3, 0xEE)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x4A, 0x55, 0x68)
DARK = RGBColor(0x1F, 0x29, 0x37)
SOFT = RGBColor(0xE8, 0xEE, 0xF4)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def build_pptx(summary: DocumentSummary, output_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    _title_slide(presentation, summary)
    _overview_slide(presentation, summary)
    _feature_table_slide(presentation, summary)
    for theme, features in _theme_groups(summary)[:5]:
        _theme_detail_slide(presentation, summary, theme, features)
    _enablement_slide(presentation, summary)
    _bullets_slide(presentation, "Next steps", (summary.next_steps or _default_next_steps(summary))[:4], "")
    _add_footers(presentation, summary)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def _blank_slide(presentation: Presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, CREAM)
    _rect(slide, 0, 0, Inches(0.16), SLIDE_HEIGHT, ORACLE_RED)
    return slide


def _title_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, NAVY)
    _rect(slide, 0, Inches(5.85), SLIDE_WIDTH, Inches(1.65), ORACLE_RED)
    product = summary.product_name or "Oracle HCM"
    release = f"Update {summary.release}" if summary.release else "Readiness update"
    _text(slide, Inches(0.7), Inches(1.35), Inches(12), Inches(0.4), "ORACLE FUSION CLOUD", 14, WHITE, bold=True)
    _text(slide, Inches(0.7), Inches(1.85), Inches(12), Inches(1.7), f"{product} {summary.release}".strip(), 40, WHITE, bold=True)
    _text(slide, Inches(0.7), Inches(3.6), Inches(12), Inches(0.6), f"{release}  ·  What's New — Feature Summary", 20, WHITE)
    _text(
        slide,
        Inches(0.7),
        Inches(4.3),
        Inches(12),
        Inches(1.1),
        summary.audience_line or "Prepared for client stakeholders.",
        14,
        RGBColor(0xD6, 0xDE, 0xE8),
    )
    _text(
        slide,
        Inches(0.7),
        Inches(6.15),
        Inches(12),
        Inches(1.1),
        f"{len(summary.features)} features  ·  {(summary.executive_summary[0] if summary.executive_summary else summary.title)}",
        14,
        WHITE,
    )


def _overview_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, "Overview", "Release highlights at a glance")
    themes = summary.themes[:5] or [
        ThemeSummary(title=feature.theme or "Enhancement", message=feature.business_benefit, features=[feature.title])
        for feature in summary.features[:5]
    ]
    positions = (
        (Inches(0.5), Inches(1.5)),
        (Inches(4.75), Inches(1.5)),
        (Inches(9.0), Inches(1.5)),
        (Inches(2.6), Inches(4.25)),
        (Inches(6.9), Inches(4.25)),
    )
    accents = (NAVY, TEAL, GOLD, ORACLE_RED, NAVY_MID)
    for index, (theme, origin, accent) in enumerate(zip(themes, positions, accents), start=1):
        left, top = origin
        _rect(slide, left, top, Inches(3.85), Inches(2.45), WHITE)
        _rect(slide, left, top, Inches(3.85), Inches(0.1), accent)
        _text(slide, left + Inches(0.2), top + Inches(0.2), Inches(3.45), Inches(0.35), f"{index:02d}", 12, accent, bold=True)
        _text(slide, left + Inches(0.2), top + Inches(0.55), Inches(3.45), Inches(0.7), theme.title, 16, NAVY, bold=True)
        _text(slide, left + Inches(0.2), top + Inches(1.3), Inches(3.45), Inches(0.9), _trim(theme.message, 70), 12, SLATE)


def _feature_table_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    features = summary.features or []
    for page, chunk in enumerate(_chunks(features, 10) or [[]], start=1):
        slide = _blank_slide(presentation)
        heading = "Feature summary" if len(features) <= 10 else f"Feature summary ({page})"
        _heading(slide, heading, "Impact: None = not enabled by default; Small scale = minimal process change")
        rows = len(chunk) + 1
        table_shape = slide.shapes.add_table(rows, 3, Inches(0.45), Inches(1.45), Inches(12.4), Inches(0.42 * rows + 0.2))
        table = table_shape.table
        table.columns[0].width = Inches(6.6)
        table.columns[1].width = Inches(2.3)
        table.columns[2].width = Inches(3.5)
        headers = ("Feature", "Impact", "Action to Enable")
        for col, header in enumerate(headers):
            _table_cell(table.cell(0, col), header, bold=True, fill=NAVY, color=WHITE)
        for row, feature in enumerate(chunk, start=1):
            _table_cell(table.cell(row, 0), feature.title, fill=WHITE)
            _table_cell(table.cell(row, 1), feature.impact or _impact_label(feature.enablement), fill=WHITE)
            _table_cell(table.cell(row, 2), _action_label(feature.enablement), fill=WHITE)


def _theme_groups(summary: DocumentSummary) -> list[tuple[ThemeSummary, list[FeatureSummary]]]:
    grouped: list[tuple[ThemeSummary, list[FeatureSummary]]] = []
    used: set[str] = set()
    for theme in summary.themes:
        names = {name.lower() for name in theme.features}
        items = [
            feature
            for feature in summary.features
            if feature.title not in used and (feature.title.lower() in names or feature.theme == theme.title)
        ]
        for feature in items:
            used.add(feature.title)
        if items:
            grouped.append((theme, items))
    leftover = [feature for feature in summary.features if feature.title not in used]
    if leftover:
        grouped.append(
            (
                ThemeSummary(title="Other enhancements", message="Additional changes in this update", features=[item.title for item in leftover]),
                leftover,
            )
        )
    return grouped


def _theme_detail_slide(
    presentation: Presentation,
    summary: DocumentSummary,
    theme: ThemeSummary,
    features: list[FeatureSummary],
) -> None:
    if len(features) == 1:
        _feature_deep_slide(presentation, theme, features[0])
        return
    for chunk in _chunks(features, 4) or [[]]:
        slide = _blank_slide(presentation)
        _heading(slide, theme.title, theme.message or f"{len(features)} enhancements")
        positions = (
            (Inches(0.5), Inches(1.5)),
            (Inches(6.9), Inches(1.5)),
            (Inches(0.5), Inches(4.2)),
            (Inches(6.9), Inches(4.2)),
        )
        for feature, origin in zip(chunk, positions):
            left, top = origin
            _rect(slide, left, top, Inches(5.9), Inches(2.45), WHITE)
            _rect(slide, left, top, Inches(0.12), Inches(2.45), TEAL)
            _text(slide, left + Inches(0.3), top + Inches(0.12), Inches(5.4), Inches(0.55), feature.title, 14, NAVY, bold=True)
            bullets = feature.details or feature.whats_new or [feature.business_benefit]
            body = "\n".join(f"• {_trim(item, 70)}" for item in bullets[:3] if item)
            _text(slide, left + Inches(0.3), top + Inches(0.7), Inches(5.4), Inches(1.25), body, 12, DARK)
            _text(slide, left + Inches(0.3), top + Inches(1.95), Inches(5.4), Inches(0.35), _action_label(feature.enablement), 11, TEAL, bold=True)


def _feature_deep_slide(presentation: Presentation, theme: ThemeSummary, feature: FeatureSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, theme.title, feature.title)
    _rect(slide, Inches(0.5), Inches(1.5), Inches(6.1), Inches(3.9), WHITE)
    _text(slide, Inches(0.7), Inches(1.65), Inches(5.7), Inches(0.4), "How it works", 14, TEAL, bold=True)
    bullets = feature.details or feature.whats_new or [feature.business_benefit]
    _text(
        slide,
        Inches(0.7),
        Inches(2.15),
        Inches(5.7),
        Inches(3.0),
        "\n\n".join(f"▸  {_trim(item, 85)}" for item in bullets[:5] if item),
        13,
        DARK,
    )
    _rect(slide, Inches(6.8), Inches(1.5), Inches(6.0), Inches(3.9), WHITE)
    _text(slide, Inches(7.0), Inches(1.65), Inches(5.6), Inches(0.4), "Setup and constraints", 14, ORACLE_RED, bold=True)
    right = feature.profile_options[:4] or feature.actions[:3] or [feature.enablement]
    right_text = "\n\n".join(f"• {_trim(item, 70)}" for item in right if item)
    if feature.business_benefit:
        right_text = f"{_trim(feature.business_benefit, 90)}\n\n{right_text}"
    _text(slide, Inches(7.0), Inches(2.15), Inches(5.6), Inches(3.0), right_text, 13, DARK)
    takeaway = feature.takeaway or "Review Steps to Enable in the Word document."
    footer = f"{feature.number}  |  {_action_label(feature.enablement)}  |  {takeaway}" if feature.number else takeaway
    _rect(slide, Inches(0.5), Inches(5.55), Inches(12.3), Inches(1.15), WHITE)
    _text(slide, Inches(0.7), Inches(5.7), Inches(12.0), Inches(0.85), f"Key takeaway  {_trim(footer, 160)}", 13, NAVY)


def _enablement_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, "Enablement", "Common profile options and access")
    profiles: list[str] = []
    for feature in summary.features:
        for option in feature.profile_options:
            if option not in profiles:
                profiles.append(option)
    cards = [(option, "Profile option") for option in profiles[:4]]
    if not cards:
        cards = [
            (str(len(summary.features)), "Features"),
            (summary.release or "—", "Release"),
            ("Opt-in", "Most items"),
            ("Word file", "Full steps"),
        ]
    width = Inches(2.85)
    left = Inches(0.5)
    for index, (value, label) in enumerate(cards[:4]):
        accent = (NAVY, TEAL, GOLD, ORACLE_RED)[index]
        _rect(slide, left, Inches(1.55), width, Inches(2.3), WHITE)
        _rect(slide, left, Inches(1.55), width, Inches(0.1), accent)
        _text(slide, left + Inches(0.12), Inches(1.8), width - Inches(0.24), Inches(1.2), str(value), 12 if len(str(value)) > 18 else 14, accent, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, left + Inches(0.12), Inches(3.1), width - Inches(0.24), Inches(0.5), label, 12, SLATE, align=PP_ALIGN.CENTER)
        left += width + Inches(0.22)
    takeaway = (
        "Most features are opt-in. Enable the relevant profile options, then complete feature-specific setup."
        if profiles
        else (summary.actions[0] if summary.actions else "Review Steps to Enable in the Word document.")
    )
    _rect(slide, Inches(0.5), Inches(4.15), Inches(12.3), Inches(2.15), WHITE)
    _text(slide, Inches(0.75), Inches(4.35), Inches(11.8), Inches(0.35), "Key takeaway", 14, ORACLE_RED, bold=True)
    _text(slide, Inches(0.75), Inches(4.8), Inches(11.8), Inches(1.2), _trim(takeaway, 200), 15, DARK)


def _impact_label(enablement: str) -> str:
    value = (enablement or "").lower()
    if value.startswith("auto"):
        return "Small scale"
    if "opt" in value or "setup" in value:
        return "None"
    return "Small scale"


def _action_label(enablement: str) -> str:
    value = (enablement or "").lower()
    if "potential" in value or value.startswith("auto") or "no setup" in value or "review" in value:
        return "Potential Setup"
    if "setup required" in value or value.startswith("setup") or "opt" in value:
        return "Setup Required"
    return enablement or "Potential Setup"


def _table_cell(cell, text: str, bold: bool = False, fill=None, color=DARK) -> None:
    cell.text = text or ""
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(11)
        paragraph.font.name = "Calibri"
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color


def _default_next_steps(summary: DocumentSummary) -> list[str]:
    return [
        "Review each feature against current processes",
        "Enable required profile options in a test environment",
        f"Train administrators on {summary.product_name or 'the'} changes",
        "Use the Word file for detailed setup steps",
    ]


def _agenda_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    items = [
        f"Release snapshot for {summary.product_name or 'this product'}",
        "Why this update matters to the client",
        "Themes the project team should align on",
        "Feature walkthrough: change, value, and client action",
        "Recommended actions, discussion, and next steps",
    ]
    _bullets_slide(presentation, "Agenda", items, "Client workshop flow")


def _snapshot_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, "Release snapshot", "Numbers to open the client conversation")
    cards = summary.snapshot or [
        (str(summary.topic_count), "Topics extracted"),
        (str(len(summary.features)), "Features"),
        (str(len(summary.themes)), "Themes"),
        (summary.release or "—", "Release"),
    ]
    width = Inches(2.85)
    gap = Inches(0.22)
    left = Inches(0.5)
    top = Inches(1.7)
    colors = (NAVY, TEAL, GOLD, ORACLE_RED)
    for index, (value, label) in enumerate(cards[:4]):
        _rect(slide, left, top, width, Inches(2.3), WHITE)
        _rect(slide, left, top, width, Inches(0.12), colors[index])
        _text(slide, left + Inches(0.2), top + Inches(0.45), width - Inches(0.4), Inches(0.9), value, 36, colors[index], bold=True, align=PP_ALIGN.CENTER)
        _text(slide, left + Inches(0.2), top + Inches(1.4), width - Inches(0.4), Inches(0.65), label, 14, SLATE, align=PP_ALIGN.CENTER)
        left = left + width + gap
    _text(
        slide,
        Inches(0.55),
        Inches(4.3),
        Inches(12.2),
        Inches(2.3),
        "\n".join(f"• {point}" for point in (summary.executive_summary[:3] or summary.why_it_matters[:3])),
        16,
        DARK,
    )


def _theme_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, "Themes to discuss with the client", "Group the conversation before walking features")
    themes = summary.themes[:4]
    positions = (
        (Inches(0.5), Inches(1.55)),
        (Inches(6.9), Inches(1.55)),
        (Inches(0.5), Inches(4.25)),
        (Inches(6.9), Inches(4.25)),
    )
    accents = (NAVY, TEAL, GOLD, ORACLE_RED)
    for theme, origin, accent in zip(themes, positions, accents):
        left, top = origin
        _rect(slide, left, top, Inches(5.9), Inches(2.45), WHITE)
        _rect(slide, left, top, Inches(0.12), Inches(2.45), accent)
        names = ", ".join(theme.features[:3])
        if len(theme.features) > 3:
            names += f" +{len(theme.features) - 3} more"
        _text(slide, left + Inches(0.35), top + Inches(0.15), Inches(5.3), Inches(0.45), theme.title, 18, accent, bold=True)
        _text(slide, left + Inches(0.35), top + Inches(0.65), Inches(5.3), Inches(0.9), theme.message, 13, DARK)
        _text(slide, left + Inches(0.35), top + Inches(1.6), Inches(5.3), Inches(0.65), names, 12, SLATE)


def _feature_slide(presentation: Presentation, summary: DocumentSummary, feature: FeatureSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, feature.title, f"{feature.number}  ·  {feature.theme or feature.section or summary.product_name}")
    _pill(slide, Inches(10.15), Inches(0.42), Inches(2.7), Inches(0.42), feature.enablement or "Review recommended")

    columns = (
        ("What's changing", feature.whats_new or feature.bullets, NAVY),
        ("Business value", [feature.business_benefit or feature.client_impact or "Confirm value with the functional owner."], TEAL),
        (
            "What the client needs to do",
            feature.actions or [f"{feature.enablement}. Assign an owner and test scenario."],
            ORACLE_RED,
        ),
    )
    left = Inches(0.5)
    width = Inches(4.0)
    for title, items, accent in columns:
        _rect(slide, left, Inches(1.55), width, Inches(4.0), WHITE)
        _rect(slide, left, Inches(1.55), width, Inches(0.12), accent)
        _text(slide, left + Inches(0.2), Inches(1.8), width - Inches(0.4), Inches(0.4), title, 15, accent, bold=True)
        body = "\n".join(f"• {_trim(item, 150)}" for item in items[:4] if item)
        _text(slide, left + Inches(0.2), Inches(2.3), width - Inches(0.4), Inches(3.0), body, 13, DARK)
        left += width + Inches(0.16)

    speak = feature.talking_points[0] if feature.talking_points else feature.client_impact
    if speak:
        _rect(slide, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.05), WHITE)
        _text(slide, Inches(0.7), Inches(5.82), Inches(12.0), Inches(0.8), f"Say this: {_trim(speak, 220)}", 14, NAVY)


def _feature_catalog_slide(presentation: Presentation, summary: DocumentSummary, features: list[FeatureSummary]) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, f"{summary.product_name} key features", "One line each. Full setup is in the Word file.")
    positions = (
        (Inches(0.5), Inches(1.5)),
        (Inches(6.9), Inches(1.5)),
        (Inches(0.5), Inches(4.2)),
        (Inches(6.9), Inches(4.2)),
    )
    for feature, origin in zip(features, positions):
        left, top = origin
        _rect(slide, left, top, Inches(5.9), Inches(2.45), WHITE)
        _rect(slide, left, top, Inches(0.12), Inches(2.45), NAVY_MID)
        _text(slide, left + Inches(0.3), top + Inches(0.12), Inches(5.4), Inches(0.7), feature.title, 15, NAVY, bold=True)
        _text(slide, left + Inches(0.3), top + Inches(0.85), Inches(5.4), Inches(0.9), _trim(feature.business_benefit or (feature.whats_new[0] if feature.whats_new else ""), 90), 13, DARK)
        _text(slide, left + Inches(0.3), top + Inches(1.85), Inches(5.4), Inches(0.4), feature.enablement, 12, TEAL, bold=True)


def _section_divider(presentation: Presentation, title: str, subtitle: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, NAVY)
    _rect(slide, 0, 0, Inches(0.16), SLIDE_HEIGHT, ORACLE_RED)
    _text(slide, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.2), title, 36, WHITE, bold=True)
    _text(slide, Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.8), subtitle, 18, RGBColor(0xD6, 0xDE, 0xE8))


def _closing_slide(presentation: Presentation, summary: DocumentSummary) -> None:
    slide = _blank_slide(presentation)
    _heading(slide, "Thank you", "Questions and decisions")
    _text(
        slide,
        Inches(0.55),
        Inches(1.7),
        Inches(12.2),
        Inches(1.2),
        f"Use the Word document for full feature text and setup. This deck is the short summary only.",
        18,
        DARK,
    )
    _rect(slide, Inches(0.5), Inches(3.2), Inches(12.3), Inches(2.6), WHITE)
    _text(slide, Inches(0.75), Inches(3.4), Inches(11.8), Inches(0.4), "Source and detailed setup", 16, ORACLE_RED, bold=True)
    _text(
        slide,
        Inches(0.75),
        Inches(3.9),
        Inches(11.8),
        Inches(1.6),
        f"Word document: full tree, headings, and enablement steps.\n{summary.source_url}",
        14,
        SLATE,
    )


def _bullets_slide(presentation: Presentation, title: str, bullets: list[str], subtitle: str = "") -> None:
    for page, chunk in enumerate(_chunks(bullets, 5) or [[]], start=1):
        slide = _blank_slide(presentation)
        heading = title if len(bullets) <= 5 else f"{title} ({page})"
        _heading(slide, heading, subtitle)
        body = "\n\n".join(f"•  {_trim(item, 90)}" for item in chunk)
        _text(slide, Inches(0.65), Inches(1.55), Inches(12.1), Inches(5.2), body, 18, DARK)


def _heading(slide, title: str, subtitle: str = "") -> None:
    _text(slide, Inches(0.5), Inches(0.28), Inches(9.4), Inches(0.55), _trim(title, 70), 24, NAVY, bold=True)
    if subtitle:
        _text(slide, Inches(0.5), Inches(0.82), Inches(9.4), Inches(0.4), _trim(subtitle, 110), 13, SLATE)


def _pill(slide, left, top, width, height, text: str) -> None:
    _rect(slide, left, top, width, height, SOFT)
    _text(slide, left, top + Inches(0.05), width, height - Inches(0.05), _trim(text, 28), 11, NAVY, bold=True, align=PP_ALIGN.CENTER)


def _add_footers(presentation: Presentation, summary: DocumentSummary) -> None:
    label = f"{summary.product_name or 'HCM'} {summary.release}  |  Client briefing".strip()
    for index, slide in enumerate(presentation.slides, start=1):
        if index == 1:
            continue
        fill = slide.shapes[0].fill
        try:
            is_navy = fill.fore_color.rgb == NAVY
        except Exception:
            is_navy = False
        if is_navy:
            continue
        _text(slide, Inches(0.5), Inches(7.12), Inches(10.5), Inches(0.28), label, 10, SLATE)
        _text(slide, Inches(11.6), Inches(7.12), Inches(1.2), Inches(0.28), str(index), 10, SLATE, align=PP_ALIGN.RIGHT)


def _rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _text(slide, left, top, width, height, text: str, size: int, color: RGBColor, bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    paragraph = tf.paragraphs[0]
    paragraph.text = text or ""
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    paragraph.font.name = "Calibri"
    paragraph.alignment = align
    tf.paragraphs[0].space_after = Pt(0)
    return box


def _chunks(items: list, size: int) -> list[list]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def _trim(text: str, limit: int) -> str:
    text = " ".join((text or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rsplit(" ", 1)[0] + "…"
