from __future__ import annotations

import base64
import os
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from xml.etree import ElementTree

import requests
from openpyxl import load_workbook

from .catalog import discover_readiness_links, match_module_name
from .models import ImplementedModule

DEFAULT_BASE_URL = "https://ehzq-test.fa.us2.oraclecloud.com"
DEFAULT_REPORT_PATH = "/Custom/Module Implemented Report.xdo"
DEFAULT_REPORT_FILE = "Module Implemented Report_Module Implemented.xlsx"
MODULE_COLUMN = "MODULE_NAME"


def load_implemented_modules(
    output_dir: Path | None = None,
    username: str | None = None,
    password: str | None = None,
) -> tuple[list[ImplementedModule], Path]:
    """Download the Fusion Module Implemented report and map MODULE_NAME to What's New books."""
    output_dir = output_dir or Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / DEFAULT_REPORT_FILE
    workbook_bytes = download_module_report(username=username, password=password)
    report_path.write_bytes(workbook_bytes)
    names = read_module_names(workbook_bytes)
    catalog = discover_readiness_links(category=None)
    modules = [
        ImplementedModule(module_name=name, catalog_item=match_module_name(name, catalog))
        for name in names
    ]
    if not modules:
        raise RuntimeError(f"No {MODULE_COLUMN} values were found in {DEFAULT_REPORT_FILE}.")
    return modules, report_path


def download_module_report(username: str | None = None, password: str | None = None) -> bytes:
    base_url = (os.getenv("FUSION_BASE_URL") or DEFAULT_BASE_URL).rstrip("/")
    username = (username or os.getenv("FUSION_USERNAME") or "").strip()
    password = (password or os.getenv("FUSION_PASSWORD") or "").strip()
    report_path = os.getenv("FUSION_BIP_REPORT_PATH") or DEFAULT_REPORT_PATH
    if not username or username.startswith("http"):
        raise RuntimeError(
            "Set FUSION_USERNAME in .env to the Fusion login user (for example john.doe), "
            "not the environment URL."
        )
    if not password:
        raise RuntimeError("Set FUSION_PASSWORD in .env to download the Module Implemented report.")

    session = requests.Session()
    session.auth = (username, password)
    session.headers.update(
        {
            "User-Agent": "HCM-Readiness-Extractor/1.0",
            "Accept": "*/*",
        }
    )
    errors: list[str] = []
    for downloader in (_run_soap_report, _run_public_report, _run_rest_report, _run_xdo_export):
        try:
            content = downloader(session, base_url, report_path, username, password)
        except Exception as exc:
            message = str(exc)
            if "invalid username or password" in message.lower():
                raise RuntimeError(
                    "Fusion rejected the BI Publisher login. Use the Fusion username "
                    "(not the https://ehzq-test... URL) and password in .env or the form."
                ) from exc
            errors.append(f"{downloader.__name__}: {exc}")
            continue
        if _looks_like_xlsx(content):
            return content
        errors.append(f"{downloader.__name__}: response was not an Excel file")
    detail = "; ".join(errors[-3:]) if errors else "no download method succeeded"
    raise RuntimeError(f"Could not download the Module Implemented report. {detail}")


def read_module_names(workbook_bytes: bytes) -> list[str]:
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            names = _module_names_from_sheet(sheet)
            if names:
                return names
    finally:
        workbook.close()
    return []


def _module_names_from_sheet(sheet) -> list[str]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    column_index = None
    for index, row in enumerate(rows[:20]):
        found = _module_column_index(row or ())
        if found is not None:
            header_index = index
            column_index = found
            break
    if header_index is None or column_index is None:
        return []
    seen: set[str] = set()
    names: list[str] = []
    for row in rows[header_index + 1 :]:
        if not row or column_index >= len(row):
            continue
        value = str(row[column_index] or "").strip()
        key = " ".join(value.lower().split())
        if not value or key in seen:
            continue
        seen.add(key)
        names.append(value)
    return names


def _module_column_index(header) -> int | None:
    for index, cell in enumerate(header):
        label = str(cell or "").strip().upper().replace(" ", "_")
        if label == MODULE_COLUMN:
            return index
    return None


def _run_soap_report(
    session: requests.Session,
    base_url: str,
    report_path: str,
    username: str,
    password: str,
) -> bytes:
    endpoint = f"{base_url}/xmlpserver/services/v2/ReportService"
    body = _soap_envelope(
        "http://xmlns.oracle.com/oxp/service/v2",
        report_path,
        username,
        password,
    )
    return _post_soap(session, endpoint, body)


def _run_public_report(
    session: requests.Session,
    base_url: str,
    report_path: str,
    username: str,
    password: str,
) -> bytes:
    endpoint = f"{base_url}/xmlpserver/services/PublicReportService"
    body = _soap_envelope(
        "http://xmlns.oracle.com/oxp/service/PublicReportService",
        report_path,
        username,
        password,
    )
    return _post_soap(session, endpoint, body)


def _soap_envelope(namespace: str, report_path: str, username: str, password: str) -> str:
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<soapenv:Envelope xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/" xmlns:pub="{namespace}">
  <soapenv:Body>
    <pub:runReport>
      <pub:reportRequest>
        <pub:attributeFormat>xlsx</pub:attributeFormat>
        <pub:attributeLocale>en-US</pub:attributeLocale>
        <pub:reportAbsolutePath>{report_path}</pub:reportAbsolutePath>
        <pub:sizeOfDataChunkDownload>-1</pub:sizeOfDataChunkDownload>
      </pub:reportRequest>
      <pub:userID>{username}</pub:userID>
      <pub:password>{password}</pub:password>
    </pub:runReport>
  </soapenv:Body>
</soapenv:Envelope>
"""


def _post_soap(session: requests.Session, endpoint: str, body: str) -> bytes:
    response = session.post(
        endpoint,
        data=body.encode("utf-8"),
        headers={
            "Content-Type": "text/xml; charset=utf-8",
            "SOAPAction": "",
        },
        timeout=180,
    )
    if "invalid username or password" in response.text.lower():
        raise RuntimeError("invalid username or password")
    response.raise_for_status()
    return _report_bytes_from_soap(response.text)


def _run_rest_report(session: requests.Session, base_url: str, report_path: str, username: str, password: str) -> bytes:
    encoded = quote(report_path.lstrip("/").replace("/", "%2F"), safe="%")
    url = f"{base_url}/xmlpserver/services/rest/v1/reports/{encoded}/run"
    response = session.post(
        url,
        files={"ReportRequest": (None, '{"attributeFormat":"xlsx","byPassCache":true}', "application/json")},
        timeout=180,
    )
    response.raise_for_status()
    if _looks_like_xlsx(response.content):
        return response.content
    return _xlsx_from_multipart(response.content)


def _run_xdo_export(session: requests.Session, base_url: str, report_path: str, username: str, password: str) -> bytes:
    encoded_path = quote(report_path, safe="/")
    urls = (
        f"{base_url}/xmlpserver{encoded_path}?_xpt=1&_xf=xlsx",
        f"{base_url}/xmlpserver/servlet/xdo?_xdo={quote(report_path)}&_xpt=1&_xf=xlsx",
    )
    last_error = "no export URL worked"
    for url in urls:
        response = session.get(url, timeout=180)
        if response.ok and _looks_like_xlsx(response.content):
            return response.content
        last_error = f"{response.status_code} {response.reason}"
    raise RuntimeError(last_error)


def _report_bytes_from_soap(xml_text: str) -> bytes:
    match = re.search(r"<(?:[\w-]+:)?reportBytes>([^<]+)</(?:[\w-]+:)?reportBytes>", xml_text)
    if match:
        return base64.b64decode(match.group(1))
    root = ElementTree.fromstring(xml_text)
    for node in root.iter():
        if node.tag.endswith("reportBytes") and node.text:
            return base64.b64decode(node.text)
    raise RuntimeError("SOAP response did not include reportBytes")


def _xlsx_from_multipart(content: bytes) -> bytes:
    marker = b"PK\x03\x04"
    start = content.find(marker)
    if start < 0:
        raise RuntimeError("REST response did not include an Excel file")
    return content[start:]


def _looks_like_xlsx(content: bytes) -> bool:
    return bool(content) and content[:2] == b"PK"
