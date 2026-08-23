from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .http_client import HttpClient
from .models import ContentBlock, DocumentTree, TopicNode

ORACLE_RED = RGBColor(0xC7, 0x46, 0x34)
NAVY = RGBColor(0x1A, 0x36, 0x5D)
GRAY = RGBColor(0x55, 0x55, 0x55)
MAX_IMAGE_WIDTH = Inches(6.1)
MAX_IMAGE_HEIGHT = Inches(3.6)


def build_docx(
    tree: DocumentTree,
    output_path: Path,
    client: HttpClient | None = None,
) -> Path:
    client = client or HttpClient()
    document = Document()
    _set_styles(document)
    _add_cover(document, tree)
    _add_index(document, tree)
    for node in tree.roots:
        _write_node(document, node, client)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _set_styles(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    for name, size, color in (
        ("Heading 1", 18, ORACLE_RED),
        ("Heading 2", 15, NAVY),
        ("Heading 3", 13, NAVY),
        ("Heading 4", 12, NAVY),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def _add_cover(document: Document, tree: DocumentTree) -> None:
    eyebrow = document.add_paragraph("Oracle Cloud Applications Readiness")
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(eyebrow, ORACLE_RED, 12, bold=True)

    title = document.add_paragraph(tree.title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style = document.styles["Title"]

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic_count = sum(1 for node in tree.walk() if node.url)
    meta.add_run(
        f"Generated {datetime.now().strftime('%d %B %Y')}\n"
        f"{topic_count} topics extracted from the document tree\n"
        f"Source: {tree.source_url}"
    ).font.size = Pt(11)

    note = document.add_paragraph(
        "This document follows the original Oracle table of contents. "
        "Section numbers match the left-nav tree. Each topic includes the "
        "source page URL under the heading."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(note, GRAY, 10)
    document.add_page_break()


def _add_index(document: Document, tree: DocumentTree) -> None:
    document.add_heading("Index", level=1)
    intro = document.add_paragraph(
        "The index below mirrors the nested links under the selected What's New URL."
    )
    _set_run(intro, GRAY, 10)

    for number, title, depth in tree.numbered_index():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25 * depth)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(f"{number}  {title}")
        run.font.size = Pt(11)
        run.bold = depth <= 1
    document.add_page_break()


def _write_node(document: Document, node: TopicNode, client: HttpClient) -> None:
    heading_level = min(node.depth + 1, 4)
    document.add_heading(f"{node.number}  {node.title}", level=heading_level)

    if node.url:
        link = document.add_paragraph()
        _add_hyperlink(link, node.url, node.url)
        _set_run(link, GRAY, 9)

    if node.error:
        warning = document.add_paragraph(f"Content could not be extracted: {node.error}")
        _set_run(warning, ORACLE_RED, 11)

    if node.content:
        for block in node.content.blocks:
            _write_block(document, block, client)

    for child in node.children:
        _write_node(document, child, client)


def _write_block(document: Document, block: ContentBlock, client: HttpClient) -> None:
    if block.kind == "heading":
        level = min(max(block.level, 2), 4)
        document.add_heading(block.text, level=level)
        return

    if block.kind == "paragraph":
        document.add_paragraph(block.text)
        return

    if block.kind == "list":
        for item in block.items:
            document.add_paragraph(item, style="List Bullet")
        return

    if block.kind == "table" and block.rows:
        _write_table(document, block.rows)
        return

    if block.kind == "image":
        _write_image(document, block, client)


def _write_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            cell.text = row[col_index] if col_index < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
    document.add_paragraph()


def _write_image(document: Document, block: ContentBlock, client: HttpClient) -> None:
    if block.image_url:
        try:
            image_bytes = client.get_bytes(block.image_url)
            width, height = _fit_image(image_bytes)
            document.add_picture(io.BytesIO(image_bytes), width=width, height=height)
        except Exception:
            fallback = document.add_paragraph(f"[Image not downloaded] {block.image_url}")
            _set_run(fallback, GRAY, 9)
    if block.caption:
        caption = document.add_paragraph(block.caption)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(caption, GRAY, 9, italic=True)


def _fit_image(image_bytes: bytes) -> tuple:
    with Image.open(io.BytesIO(image_bytes)) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        return MAX_IMAGE_WIDTH, None
    aspect = height_px / width_px
    width = MAX_IMAGE_WIDTH
    height = Inches(width.inches * aspect)
    if height > MAX_IMAGE_HEIGHT:
        height = MAX_IMAGE_HEIGHT
        width = Inches(height.inches / aspect)
    return width, height


def _set_run(paragraph, color: RGBColor, size: int, bold: bool = False, italic: bool = False) -> None:
    if not paragraph.runs:
        return
    run = paragraph.runs[0]
    run.font.color.rgb = color
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A365D")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(color)
    r_pr.append(underline)
    r_pr.append(size)
    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def safe_filename(title: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("_")
    return (cleaned or "oracle_readiness")[:80]
